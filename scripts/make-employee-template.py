"""Generate the styled bulk-employee-creation Excel template.

Output: public/employee-upload-template.xlsx
Run:    python3 scripts/make-employee-template.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BRAND = "2E4566"          # Tasknet brand
BRAND_DARK = "24374F"
AMBER_TINT = "FFF7E6"     # example-row highlight
GREY_BORDER = "D8DEE7"
TEXT = "1F2937"

HEADERS = [
    "Employee Code *",
    "Employee Name *",
    "Email ID *",
    "Process *",
    "Designation *",
    "L1 Manager Name",
    "L1 Manager Email ID",
]
EXAMPLE = [
    "EXAMPLE",
    "Avinash Sharma",
    "avinash@digitide.com",
    "Customer Support",
    "Executive",
    "Priya Nair",
    "priya.nair@digitide.com",
]
WIDTHS = [16, 24, 30, 20, 22, 20, 30]

thin = Side(style="thin", color=GREY_BORDER)
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)


def build_employees_sheet(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Employees"

    # Header row
    for col, title in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=title)
        c.font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BRAND)
        c.alignment = center
        c.border = border
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    # Example row (parser ignores Employee Code "EXAMPLE")
    for col, val in enumerate(EXAMPLE, start=1):
        c = ws.cell(row=2, column=col, value=val)
        c.font = Font(name="Calibri", italic=True, size=10, color="6B7280")
        c.fill = PatternFill("solid", fgColor=AMBER_TINT)
        c.alignment = left
        c.border = border
    ws.row_dimensions[2].height = 20

    # Pre-drawn entry grid for the next 40 rows
    for row in range(3, 43):
        for col in range(1, 8):
            c = ws.cell(row=row, column=col)
            c.border = border
            c.alignment = left
            c.font = Font(name="Calibri", size=11, color=TEXT)

    for col, width in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def build_instructions_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Instructions")
    ws.sheet_view.showGridLines = False
    for col, width in enumerate([3, 26, 12, 78], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    row = 2
    c = ws.cell(row=row, column=2, value="Tasknet — Bulk Employee ID Creation")
    c.font = Font(name="Calibri", bold=True, size=16, color=BRAND_DARK)
    row += 1
    c = ws.cell(row=row, column=2, value="Admin-only upload template · one employee per row")
    c.font = Font(name="Calibri", size=10, color="6B7280")

    row += 2
    c = ws.cell(row=row, column=2, value="How to use")
    c.font = Font(name="Calibri", bold=True, size=12, color=BRAND)
    steps = [
        "1.  Open the Employees sheet and fill employee details — one employee per row, starting at row 3.",
        "2.  Replace or delete the grey EXAMPLE row (rows whose Employee Code is \"EXAMPLE\" are ignored).",
        "3.  Keep the header row exactly as it is — do not rename, reorder or delete columns.",
        "4.  Save the file and upload it in Tasknet → Admin → Bulk Create via Excel.",
    ]
    for s in steps:
        row += 1
        ws.cell(row=row, column=2, value=s).font = Font(name="Calibri", size=10.5, color=TEXT)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)

    row += 2
    c = ws.cell(row=row, column=2, value="Columns")
    c.font = Font(name="Calibri", bold=True, size=12, color=BRAND)
    row += 1
    for title, offset in [("Column", 0), ("Required", 1), ("What to enter", 2)]:
        c = ws.cell(row=row, column=2 + offset, value=title)
        c.font = Font(name="Calibri", bold=True, size=10.5, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BRAND)
        c.border = border
        c.alignment = center if offset == 1 else left
    rows = [
        ("Employee Code *", "Yes", "Unique login code, e.g. EMP101"),
        ("Employee Name *", "Yes", "Full name, e.g. Karan Mehta"),
        ("Email ID *", "Yes", "Unique work email, e.g. karan.mehta@digitide.com"),
        ("Process *", "Yes", "Team / process, e.g. Customer Support"),
        ("Designation *", "Yes", "e.g. Executive / TL / AM / DM"),
        ("L1 Manager Name", "No", "Manager's full name"),
        ("L1 Manager Email ID", "No", "Must match an existing employee's Email ID (or another row in this file) to auto-link the hierarchy"),
    ]
    for name, req, desc in rows:
        row += 1
        a = ws.cell(row=row, column=2, value=name)
        b = ws.cell(row=row, column=3, value=req)
        d = ws.cell(row=row, column=4, value=desc)
        for c_ in (a, b, d):
            c_.border = border
            c_.font = Font(name="Calibri", size=10.5, color=TEXT)
            c_.alignment = Alignment(horizontal="center" if c_ is b else "left", vertical="center", wrap_text=True)
        a.font = Font(name="Calibri", size=10.5, bold=True, color=TEXT)
        b.font = Font(name="Calibri", size=10.5, bold=req == "Yes", color="B45309" if req == "Yes" else "6B7280")

    row += 2
    c = ws.cell(row=row, column=2, value="Good to know")
    c.font = Font(name="Calibri", bold=True, size=12, color=BRAND)
    notes = [
        "•  Every created ID starts with the default password Digitide@123 and must set a new password at first login.",
        "•  Duplicate Employee Codes or Email IDs (inside the file, or already in the system) are rejected — every other row is still created.",
        "•  The admin can view or copy each user's password from the All Users table.",
        "•  Leave the L1 Manager columns blank for top-level employees.",
        "•  Up to 500 employees per upload; rows are processed top to bottom.",
    ]
    for n in notes:
        row += 1
        ws.cell(row=row, column=2, value=n).font = Font(name="Calibri", size=10.5, color=TEXT)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=4)


wb = Workbook()
build_employees_sheet(wb)
build_instructions_sheet(wb)
wb.properties.creator = "Z.ai"
wb.properties.title = "Tasknet — Employee Upload Template"

OUT = "/home/z/my-project/public/employee-upload-template.xlsx"
wb.save(OUT)
print("saved", OUT)
